"""Independent source readers. Standard library csv (and openpyxl for .xlsx).

Deliberately does NOT use pandas, so a parsing quirk in the analysis engine
cannot quietly reproduce itself here.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
from datetime import date, datetime
from pathlib import Path


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def locate(folder: Path, stem: str) -> Path | None:
    for ext in (".csv", ".xlsx"):
        candidate = folder / f"{stem}{ext}"
        if candidate.is_file():
            return candidate
    return None


SCAN_LIMIT = 20


def header_key(name) -> str:
    text = "" if name is None else str(name).strip().lower()
    return re.sub(r"[\s\-]+", "_", text)


def _empty(cell) -> bool:
    return cell is None or str(cell).strip() == ""


def _grids(path: Path):
    """Yield candidate grids (list of rows) in search order."""
    if path.suffix == ".csv":
        text = path.read_text(encoding="utf-8-sig")
        for delim in (",", ";", "\t"):
            yield list(csv.reader(io.StringIO(text, newline=""), delimiter=delim))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        grids = [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
        wb.close()
        yield from grids


def raw_table(path: Path, required: list[str]) -> tuple[list[str], list[list], int]:
    """Return (normalized header, body rows, 1-based header line). Header = first row
    within SCAN_LIMIT that names every required column. Fully blank rows and rows that
    repeat the header are dropped."""
    for rows in _grids(path):
        for i, row in enumerate(rows[:SCAN_LIMIT]):
            names = [header_key(c) for c in row]
            if set(required) <= set(names):
                pos = {n: names.index(n) for n in required}
                body = []
                for r in rows[i + 1:]:
                    if all(_empty(c) for c in r):
                        continue
                    if all(pos[n] < len(r) and header_key(r[pos[n]]) == n for n in required):
                        continue
                    body.append(r)
                return names, body, i + 1
    return [], [], 0


def _clean_cat(text: str) -> str:
    return re.sub(r"[\s_\-]+", " ", text.lower()).strip()


def parse_cell(value, typ: str):
    """Return (parsed, ok). Dates are handled by parse_date."""
    if typ in ("str", "id", "cat", "cat?", "key"):
        text = "" if value is None else str(value).strip()
        if typ == "key":
            return (header_key(text), True) if text else (None, False)
        if typ == "cat?":
            return (_clean_cat(text) or "(blank)"), True
        if not text:
            return None, False
        if typ == "id":
            return text.upper(), True
        if typ == "cat":
            return _clean_cat(text), True
        return text, True
    if typ == "float":
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            num = float(value)
        else:
            text = str(value).strip()
            if re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+(\.\d+)?", text):
                text = text.replace(",", "")
            try:
                num = float(text)
            except (TypeError, ValueError):
                return None, False
        return (num, True) if math.isfinite(num) else (None, False)
    raise ValueError(f"unknown type {typ}")


def parse_date(value, fmt: str):
    if isinstance(value, datetime):
        return value.date().isoformat(), True
    if isinstance(value, date):
        return value.isoformat(), True
    text = "" if value is None else str(value).strip()
    try:
        return datetime.strptime(text, fmt).date().isoformat(), True
    except ValueError:
        return None, False


def read_typed(path: Path, columns: dict[str, str], total_pattern: str, date_formats: list[str]):
    """Return dict with header, header_line, rows (typed dicts), cells, failures, total_rows."""
    required = list(columns)
    header, body, header_line = raw_table(path, required)
    if not header:
        return {"header": [], "header_line": 0, "rows": [], "cells": 0, "total_rows": 0,
                "failures": [{"line": 0, "column": "*", "value": "no header row with all required columns"}]}
    total_re = re.compile(total_pattern, re.I)
    data = []
    total_rows = 0
    for offset, raw in enumerate(body):
        first = "" if not raw or raw[0] is None else str(raw[0])
        if total_re.match(first):
            total_rows += 1
            continue
        data.append((offset, raw))
    idx = {name: header.index(name) for name in columns}

    def cell(raw, name):
        i = idx[name]
        return raw[i] if i < len(raw) else None

    # choose one date format per column: the first that parses every value
    chosen = {}
    for name, typ in columns.items():
        if typ == "date":
            chosen[name] = date_formats[0]
            for fmt in date_formats:
                if all(parse_date(cell(raw, name), fmt)[1] for _, raw in data):
                    chosen[name] = fmt
                    break

    rows, failures, cells = [], [], 0
    for offset, raw in data:
        rec = {}
        for name, typ in columns.items():
            cells += 1
            value = cell(raw, name)
            val, ok = parse_date(value, chosen[name]) if typ == "date" else parse_cell(value, typ)
            if not ok:
                failures.append({"line": header_line + 1 + offset, "column": name,
                                 "value": None if value is None else str(value)})
            rec[name] = val
        rows.append(rec)
    return {"header": header, "header_line": header_line, "rows": rows, "cells": cells,
            "failures": failures, "total_rows": total_rows}
