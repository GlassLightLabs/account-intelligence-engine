"""Validation gate 2: check finished documents before they are released.

Every number in a customer document must either carry a claim marker that
matches an independent recompute, or be explicitly marked as a literal.
"""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path

from .context import Account
from .readers import file_sha256
from .report import Report

CLAIM = re.compile(r"(?P<val>[+-]?\d[\d,]*(?:\.\d+)?)(?P<unit>%| pts)?<!--claim:(?P<ref>[^>]*?)-->")
LITERAL = re.compile(r"(?P<val>[+-]?\d[\d,]*(?:\.\d+)?)<!--literal-->")
COMMENT = re.compile(r"<!--.*?-->", re.S)
BARE_NUMBER = re.compile(r"(?<![\w\-/.:@\[])[+-]?\d[\d,]*(?:\.\d+)?%?(?![\w\-/:\]])")
MONTH_DATE = re.compile(r"\b(Jan(uary)?|Feb(ruary)?|Mar(ch)?|Apr(il)?|May|June?|July?|Aug(ust)?|"
                        r"Sep(t(ember)?)?|Oct(ober)?|Nov(ember)?|Dec(ember)?)\.?\s+\d{1,2}(st|nd|rd|th)?"
                        r"(,?\s+\d{4})?\b")
ORDINAL = re.compile(r"^\s*\d+\.\s", re.M)
FACTS_SHA = re.compile(r"<!--\s*facts-sha256:([0-9a-f]{64})\s*-->")
REF = re.compile(r"^(?:(?P<site>[^@:]+)@(?P<period>\d{4}-\d{2}):)?(?P<body>.+)$")


def _displayed(val: str, unit: str | None) -> tuple[float, float]:
    """Parse a displayed number. Returns (value in fact units, rounding tolerance)."""
    num = float(val.replace(",", ""))
    decimals = len(val.split(".")[1]) if "." in val else 0
    tol = 0.5 * 10 ** -decimals + 1e-9
    if unit in ("%", " pts"):
        return num / 100, tol / 100
    return num, tol


def resolve(acct: Account, ref: str, default_site: str | None, default_period: str | None):
    """Return (value, error)."""
    m = REF.match(ref.strip())
    if not m:
        return None, "unparseable claim reference"
    site = m.group("site") or default_site
    period = m.group("period") or default_period
    body = m.group("body")
    spec = acct.spec
    if body.startswith("target:"):
        metric = body.split(":", 1)[1]
        t = spec.get("targets", {}).get(metric)
        return (t, None) if t is not None else (None, f"no target for {metric}")
    if not site or not period:
        return None, "claim has no site/period"
    if acct.compute(site, period) is None:
        return None, f"cannot recompute {site} {period}"
    if body.startswith(("delta:", "absdelta:")):
        kind, target = body.split(":", 1)
        bm = re.match(r"^(?P<name>[a-z_]+)\[(?P<key>.+)\]$", target)
        if bm:
            prev = acct.previous_period(site, period)
            cur = acct.breakdown_item(site, period, bm.group("name"), bm.group("key"))
            old = acct.breakdown_item(site, prev, bm.group("name"), bm.group("key")) if prev else None
            d = None if cur is None or old is None else cur - old
        else:
            d = acct.delta(site, period, target)
        if d is None:
            return None, f"no previous value for {body}"
        return (abs(d) if kind == "absdelta" else d), None
    if body.startswith("baseline:"):
        mean, _, _ = acct.baseline(site, period, body.split(":", 1)[1])
        return (mean, None) if mean is not None else (None, "no baseline history")
    bm = re.match(r"^(?P<name>[a-z_]+)\[(?P<key>.+)\]$", body)
    if bm:
        if bm.group("name") not in spec.get("breakdowns", {}):
            return None, f"unknown breakdown {bm.group('name')}"
        v = acct.breakdown_item(site, period, bm.group("name"), bm.group("key"))
        return (v, None) if v is not None else (None, f"{bm.group('key')} is not in {bm.group('name')}")
    if body not in spec["metrics"]:
        return None, f"unknown metric {body}"
    v = acct.value(site, period, body)
    return (v, None) if v is not None else (None, f"{body} has no value")


def check_document(acct: Account, rep: Report, path: Path, label: str,
                   default_site: str | None, default_period: str | None) -> None:
    text = path.read_text(encoding="utf-8")
    claims = list(CLAIM.finditer(text))
    internal = {n for n, b in acct.spec.get("breakdowns", {}).items() if b.get("internal")}
    for i, mt in enumerate(claims):
        ref = mt.group("ref")
        cited = re.findall(r"([a-z_]+)\[", ref)
        if any(n in internal for n in cited):
            rep.add("claims", "hard_fail", f"{label}#{i + 1}.internal",
                    f"cites {ref}, which the pack marks internal-only (for example, individual-level data)")
            continue
        expected, err = resolve(acct, ref, default_site, default_period)
        shown = mt.group("val") + (mt.group("unit") or "")
        if err:
            rep.add("claims", "hard_fail", f"{label}#{i + 1}", f"'{shown}' cites {ref}: {err}")
            continue
        val, tol = _displayed(mt.group("val"), mt.group("unit"))
        rep.expect("claims", f"{label}#{i + 1}", abs(val - expected) <= tol,
                   f"'{shown}' does not match {ref}", ref=ref, expected=expected, shown=shown)
    for mt in LITERAL.finditer(text):
        rep.add("claims", "info", f"{label}.literal", f"literal number '{mt.group('val')}' (not a data claim)")

    stripped = CLAIM.sub(" ", text)
    stripped = LITERAL.sub(" ", stripped)
    stripped = COMMENT.sub(" ", stripped)
    stripped = re.sub(r"`[^`]*`", " ", stripped)
    stripped = ORDINAL.sub(" ", stripped)
    stripped = MONTH_DATE.sub(" ", stripped)   # "August 3", "Sep 28, 2026" are dates, not data
    bare = BARE_NUMBER.findall(stripped)
    rep.expect("claims", f"{label}.unsupported_numbers", not bare,
               f"{len(bare)} number(s) have no claim marker: {bare[:10]}", level="warning")
    rep.add("claims", "info", f"{label}.claim_count", f"{len(claims)} claims checked")


def check_workbook(acct: Account, rep: Report, path: Path, site: str, period: str) -> None:
    from openpyxl import load_workbook
    tol = float(acct.spec.get("integrity", {}).get("tolerance_abs", 1e-6))
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Metrics"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    ci, cv = header.index("Metric"), header.index("Value")
    for r in rows[1:]:
        metric, value = r[ci], r[cv]
        exp = acct.value(site, period, metric)
        ok = (value is None and exp is None) or (value is not None and exp is not None and abs(value - exp) <= tol)
        rep.expect("workbook", f"Metrics.{metric}", ok, "workbook value does not match recompute",
                   expected=exp, actual=value)
    for name in acct.spec.get("breakdowns", {}):
        if name[:31] not in wb.sheetnames:
            continue
        mine = dict(acct.compute(site, period)["breakdowns"].get(name, []))
        for key, value in list(wb[name[:31]].iter_rows(values_only=True))[1:]:
            exp = mine.get(key)
            rep.expect("workbook", f"{name}[{key}]", exp is not None and abs(value - exp) <= tol,
                       "workbook breakdown value does not match recompute", expected=exp, actual=value)
    wb.close()


def check_docs(acct: Account, site: str, period: str) -> dict:
    rep = Report("docs", site, period)
    facts_path = acct.facts_path(site, period)
    v_path = acct.validation_path(site, period, "facts")
    gate1 = json.loads(v_path.read_text(encoding="utf-8")) if v_path.is_file() else None
    current_sha = file_sha256(facts_path) if facts_path.is_file() else None
    rep.expect("gate", "facts_validated", bool(gate1) and gate1["status"] == "VALIDATED",
               "facts have not passed gate 1")
    rep.expect("gate", "facts_unchanged", bool(gate1) and gate1.get("facts_sha256") == current_sha,
               "facts changed after gate 1 validation")

    out = acct.outputs_dir(site, period)
    overview = out / "site_overview.md"
    if rep.expect("gate", "site_overview_present", overview.is_file(), "site_overview.md not found"):
        m = FACTS_SHA.search(overview.read_text(encoding="utf-8"))
        rep.expect("provenance", "site_overview.facts_sha256", bool(m) and m.group(1) == current_sha,
                   "document was rendered from different facts", expected=current_sha,
                   actual=m.group(1) if m else None)
        check_document(acct, rep, overview, "site_overview", site, period)
    workbook = out / "analysis.xlsx"
    if workbook.is_file():
        check_workbook(acct, rep, workbook, site, period)
    summary = acct.root / "outputs" / "account_summary.md"
    if summary.is_file():
        check_document(acct, rep, summary, "account_summary", None, None)

    data = rep.write(acct.validation_path(site, period, "docs"),
                     acct.validation_path(site, period, "docs", "md"), facts_sha256=current_sha)
    stamp = {"released_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
             "facts_sha256": current_sha, "checks": data["summary"]["checks"]}
    if data["status"] == "VALIDATED":
        if overview.is_file():
            stamp["site_overview_sha256"] = file_sha256(overview)
            (out / "RELEASED").write_text(json.dumps(stamp, indent=2) + "\n", encoding="utf-8")
        if summary.is_file():
            (acct.root / "outputs" / "account_summary.RELEASED").write_text(
                json.dumps({**stamp, "account_summary_sha256": file_sha256(summary)}, indent=2) + "\n",
                encoding="utf-8")
    else:
        for marker in (out / "RELEASED", acct.root / "outputs" / "account_summary.RELEASED"):
            if marker.exists():
                marker.unlink()
    return data


def check_interpretation(acct: Account, site: str, period: str) -> dict:
    """Pre-check Claude's interpretation file before rendering: every claim must resolve and match,
    no internal-only data, no unsupported numbers. Cheap to run, and it catches mistakes before a
    full render and gate 2 cycle. Writes no RELEASED marker."""
    rep = Report("claims", site, period)
    path = acct.root / "briefings" / site / f"{period}.interpretation.md"
    if rep.expect("gate", "interpretation_present", path.is_file(), f"{path} not found"):
        check_document(acct, rep, path, "interpretation", site, period)
    return rep.write(acct.validation_path(site, period, "claims"),
                     acct.validation_path(site, period, "claims", "md"))
