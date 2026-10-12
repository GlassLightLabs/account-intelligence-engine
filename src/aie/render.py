"""Step 5 of the validated flow: build customer-facing documents.

Rendering is gated. It refuses to run unless the integrity engine has
VALIDATED the exact facts file being rendered (matched by SHA-256). Every
number in the output carries a claim marker so the second gate can check
the finished document against an independent recompute.
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .formatting import claim, fmt_delta, fmt_value, literal
from .investigations import site_items
from .store import AccountStore, now_iso, read_json, sha256_file

SHORT_CLAIM = re.compile(r"<!--claim:(?![^>]*@)([^>]*?)-->")


class GateError(Exception):
    pass


def _layout(spec: dict) -> dict:
    """Which metrics and breakdowns appear in documents. Comes from the pack's
    optional `report:` section; defaults to everything in spec order."""
    metrics = list(spec["metrics"])
    rep = spec.get("report", {})
    headline = rep.get("headline_metric", metrics[0])
    return {
        "headline": headline,
        "key_metrics": rep.get("key_metrics", metrics),
        "breakdowns": rep.get("breakdowns", list(spec.get("breakdowns", {}))),
        "summary_metrics": rep.get("summary_metrics", []),
        "data_notes": rep.get("data_notes", []),
    }


def _target_lines(spec: dict) -> list[str]:
    lines = []
    for m, t in spec.get("targets", {}).items():
        ms = spec["metrics"][m]
        lines.append(f"{ms['label']} target: {claim(fmt_value(t, ms['format']), f'target:{m}')}")
    return ["", *lines] if lines else []


def check_gate(store: AccountStore, site: str, period: str) -> dict:
    facts_path = store.facts_file(site, period)
    if not facts_path.exists():
        raise GateError(f"No facts for {site} {period}. Run `aie run` first.")
    v = read_json(store.validation_file(site, period, "facts"))
    if v is None:
        raise GateError(f"{site} {period} has not been validated. Run `aie-integrity facts` first.")
    if v["facts_sha256"] != sha256_file(facts_path):
        raise GateError(f"Facts for {site} {period} changed after validation. Re-run `aie-integrity facts`.")
    if v["status"] != "VALIDATED":
        raise GateError(f"Validation for {site} {period} is {v['status']}: "
                        f"{v['summary']['hard_fail']} hard failures. See the validation report.")
    return v


def _ref(site: str, period: str, key: str) -> str:
    return f"{site}@{period}:{key}"


def _signal_sentence(sig: dict, spec: dict, site: str, period: str) -> str:
    ev = sig["evidence"]
    if "breakdown" in ev:
        b = ev["breakdown"]
        fmt = spec["breakdowns"][b]["format"]
        val = claim(fmt_value(ev["value"], fmt), _ref(site, period, f"{b}[{ev['item']}]"))
        return f"'{ev['item']}' became the leading item in {spec['breakdowns'][b]['label'].lower()} ({val})."
    m = sig["metric"]
    ms = spec["metrics"][m]
    fmt = ms["format"]
    cur = claim(fmt_value(ev["current"], fmt), _ref(site, period, m))
    if "streak" in ev:
        start = claim(fmt_value(ev["start_value"], fmt), _ref(site, ev["start_period"], m))
        verb = "risen" if ev["direction"] == "up" else "fallen"
        return f"{ms['label']} has {verb} every month since {ev['start_period']}, from {start} to {cur}."
    if "previous" in ev:
        prev = claim(fmt_value(ev["previous"], fmt), _ref(site, ev["previous_period"], m))
        d = claim(fmt_delta(ev["delta"], fmt), _ref(site, period, f"delta:{m}"))
        return f"{ms['label']} moved from {prev} in {ev['previous_period']} to {cur} ({d})."
    if "target" in ev:
        tgt = claim(fmt_value(ev["target"], fmt), f"target:{m}")
        return f"{ms['label']} is {cur} against a target of {tgt}."
    if "baseline" in ev:
        base = claim(fmt_value(ev["baseline"], fmt), _ref(site, period, f"baseline:{m}"))
        return f"{ms['label']} came in at {cur}, well above the recent baseline of {base}."
    return sig["title"] + "."


def _persisting_sentence(item: dict, spec: dict, facts: dict, site: str, period: str) -> str:
    """Customer wording for a problem whose alert went quiet but has not recovered."""
    rule = spec["signals"][item["rule"]]
    if rule["type"] == "new_top_item":
        b = rule["breakdown"]
        bs = spec["breakdowns"][b]
        items = facts["breakdowns"].get(b, {}).get("items", [])
        val = next((i["value"] for i in items if i["key"] == item["reference"]), None)
        shown = claim(fmt_value(val, bs["format"]), _ref(site, period, f"{b}[{item['reference']}]"))
        return f"- '{item['reference']}' is still the leading item in {bs['label'].lower()} ({shown})."
    m = rule["metric"]
    ms = spec["metrics"][m]
    cur = claim(fmt_value(facts["metrics"][m]["value"], ms["format"]), _ref(site, period, m))
    src = item.get("reference_source")
    if not src:
        return f"- {ms['label']} has not yet returned to its earlier level ({cur})."
    ref = claim(fmt_value(item["reference"], ms["format"]), src["claim"])
    return f"- {ms['label']} has not recovered: {cur} this period versus {ref} before the issue began ({src['label']})."


def render_site(store: AccountStore, site: str, period: str) -> dict:
    validation = check_gate(store, site, period)
    spec = store.spec()
    facts = store.load_facts(site, period)
    comp = facts["comparison"]
    prev_p = comp["previous_period"]
    out = store.outputs_dir(site, period)
    out.mkdir(parents=True, exist_ok=True)
    released = out / "RELEASED"
    if released.exists():
        released.unlink()

    n_checks = literal(f"{validation['summary']['checks']:,}")
    n_fail = literal(str(validation["summary"]["hard_fail"]))
    L = [f"# Site overview: {store.site_name(site)}, {period}", "",
         f"**{facts['account']}** | Prepared {now_iso()[:10]} | "
         f"Data validated: {n_checks} independent checks, {n_fail} failures", ""]

    L += ["## Key metrics", "", "| Metric | " + period + " | " + (prev_p or "Previous") + " | Change |",
          "|---|---|---|---|"]
    layout = _layout(spec)
    for m in layout["key_metrics"]:
        ms, c = spec["metrics"][m], comp["metrics"][m]
        cur = claim(fmt_value(c["current"], ms["format"]), _ref(site, period, m))
        prev = claim(fmt_value(c["previous"], ms["format"]), _ref(site, prev_p, m)) if prev_p else "n/a"
        d = claim(fmt_delta(c["delta"], ms["format"]), _ref(site, period, f"delta:{m}")) if prev_p else "n/a"
        L.append(f"| {ms['label']} | {cur} | {prev} | {d} |")
    L += _target_lines(spec)
    L.append("")

    for b in layout["breakdowns"]:
        bd = facts["breakdowns"].get(b)
        if not bd or not bd["items"]:
            continue
        bs = spec["breakdowns"][b]
        L += [f"## {bs['label']}", "", "| | Value |", "|---|---|"]
        for it in bd["items"]:
            val = claim(fmt_value(it["value"], bs["format"]), _ref(site, period, f"{b}[{it['key']}]"))
            L.append(f"| {it['key']} | {val} |")
        L.append("")

    L += ["## What changed", ""]
    L += [f"- {_signal_sentence(s, spec, site, period)}" for s in facts["signals"]]
    L += [_persisting_sentence(i, spec, facts, site, period)
          for i in site_items(store, site, ("open",)) if i.get("persisting")]
    if L[-1] == "":
        L.append("- Performance held steady. No notable changes this period.")
    L.append("")

    L += ["## What this means", ""]
    interp = store.interpretation_file(site, period)
    has_interp = interp.exists()
    if has_interp:
        text = interp.read_text(encoding="utf-8").strip()
        text = SHORT_CLAIM.sub(lambda mt: f"<!--claim:{site}@{period}:{mt.group(1)}-->", text)
        text = re.sub(r"^#\s.*\n+", "", text)   # drop a leading H1 if Claude added one
        L += [text, ""]
    else:
        L += ["_Interpretation pending._", ""]

    L += ["## Data notes", ""]
    for m in layout["data_notes"]:
        ms = spec["metrics"][m]
        val = claim(fmt_value(facts["metrics"][m]["value"], ms["format"]), _ref(site, period, m))
        L.append(f"- {ms['label']}: {val}")
    L += [
          "- Sources: " + ", ".join(f"`{m['file']}`" for m in facts["inputs"].values()) + ".",
          "", f"<!-- facts-sha256:{validation['facts_sha256']} -->", ""]
    md_path = out / "site_overview.md"
    md_path.write_text("\n".join(L), encoding="utf-8")

    xlsx_path = out / "analysis.xlsx"
    _workbook(xlsx_path, facts, spec, validation)
    store.log_event("rendered", site=site, period=period, interpretation=has_interp)
    summary = render_account_summary(store)
    return {"site_overview": str(md_path), "workbook": str(xlsx_path),
            "account_summary": summary, "interpretation": has_interp}


def _workbook(path: Path, facts: dict, spec: dict, validation: dict) -> None:
    wb = Workbook()
    bold = Font(bold=True)
    ws = wb.active
    ws.title = "Metrics"
    ws.append(["Metric", "Label", "Value", "Previous", "Change", "Assessment", "Source file", "Rule"])
    for m, ms in spec["metrics"].items():
        c, f = facts["comparison"]["metrics"][m], facts["metrics"][m]
        src = f.get("source") or {}
        ws.append([m, ms["label"], c["current"], c["previous"], c["delta"], c["assessment"],
                   src.get("file"), src.get("rule")])
    for name, bd in facts["breakdowns"].items():
        s = wb.create_sheet(name[:31])
        s.append(["Key", "Value"])
        for it in bd["items"]:
            s.append([it["key"], it["value"]])
    s = wb.create_sheet("Signals")
    s.append(["Severity", "Kind", "Title", "Evidence"])
    for sig in facts["signals"]:
        s.append([sig["severity"], sig["kind"], sig["title"], str(sig["evidence"])])
    s = wb.create_sheet("Provenance")
    s.append(["Input", "File", "SHA-256", "Rows", "Date min", "Date max"])
    for k, m in facts["inputs"].items():
        s.append([k, m["file"], m["sha256"], m["rows"], m.get("date_min"), m.get("date_max")])
    s.append([])
    s.append(["Validation", validation["status"], f"{validation['summary']['checks']} checks",
              f"{validation['summary']['hard_fail']} failed", f"{validation['summary']['warning']} warnings"])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = bold
    wb.save(path)


def _without_date(text: str) -> str:
    return re.sub(r"^Updated .*$", "", text, flags=re.M)


def render_account_summary(store: AccountStore) -> str:
    """Rolling account view across sites. Only validated periods are included."""
    spec = store.spec()
    meta = store.meta()
    layout = _layout(spec)
    rows, pending, inv_lines = [], [], []
    for site in store.sites():
        validated = None
        for p in reversed(store.periods(site)):
            v = read_json(store.validation_file(site, p, "facts"))
            if v and v["status"] == "VALIDATED" and v["facts_sha256"] == sha256_file(store.facts_file(site, p)):
                validated = p
                break
        if validated is None:
            pending.append(site)
            continue
        f = store.load_facts(site, validated)
        c = f["comparison"]["metrics"]
        prev = f["comparison"]["previous_period"]
        hm = layout["headline"]
        hfmt = spec["metrics"][hm]["format"]
        cells = [claim(fmt_value(c[hm]["current"], hfmt), _ref(site, validated, hm)),
                 claim(fmt_delta(c[hm]["delta"], hfmt), _ref(site, validated, f"delta:{hm}")) if prev else "n/a"]
        for m in layout["summary_metrics"]:
            cells.append(claim(fmt_value(c[m]["current"], spec["metrics"][m]["format"]), _ref(site, validated, m)))
        flag_list = [s["title"] for s in f["signals"] if s["kind"] in ("risk", "anomaly")]
        flag_list += [f"not recovered: {i['title']}" for i in site_items(store, site, ("open",)) if i.get("persisting")]
        flags = ", ".join(flag_list) or "none"
        rows.append(f"| {store.site_name(site)} | {validated} | " + " | ".join(cells) + f" | {flags} |")
        for i in site_items(store, site):
            inv_lines.append(f"- **{site}**: {i['title']} ({i['status']}, opened {i['opened']})")

    cols = ["Site", "Period", spec["metrics"][layout["headline"]]["label"], "Change",
            *[spec["metrics"][m]["label"] for m in layout["summary_metrics"]], "Flags"]
    L = [f"# Account summary: {meta['name']}", "", f"Updated {now_iso()[:10]}. Validated data only.", "",
         "| " + " | ".join(cols) + " |", "|" + "---|" * len(cols), *rows]
    L += _target_lines(spec) + [""]
    L += ["## Open investigations", ""] + (inv_lines or ["- None."]) + [""]
    if pending:
        L += ["## Pending validation", "", *[f"- {s}" for s in pending], ""]
    content = "\n".join(L)

    out = store.root / "outputs" / "account_summary.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        old = out.read_text(encoding="utf-8")
        if _without_date(old) != _without_date(content):
            archive = store.root / "outputs" / "archive"
            archive.mkdir(exist_ok=True)
            shutil.copy2(out, archive / f"account_summary.{now_iso().replace(':', '')}.md")
    out.write_text(content, encoding="utf-8")
    rel = store.root / "outputs" / "account_summary.RELEASED"
    if rel.exists():
        rel.unlink()
    return str(out)
